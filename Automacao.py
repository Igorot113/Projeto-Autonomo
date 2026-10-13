import pyautogui
import time
from openpyxl import load_workbook
import pyperclip
from decimal import Decimal, ROUND_HALF_UP
import PegarPosicoes as pos
from pathlib import Path
# shift +  tab ao invez de dar um salto para frente voce da um salto para traz
PASTA_PROJETO = Path(__file__).parent
CAMINHO_PLANILHA = PASTA_PROJETO / "Planilhas" / "Comparação de Encerrantes.xlsx"
wb = load_workbook(CAMINHO_PLANILHA)
aba = wb.active

lista_bicos = []
lista_nomes = []
lista_diferenca = []
lista_bicos_nomes = []
lista_linhas_Excel_Segunda_Forma = []
execelInicio = input("Qual é o numero da linha inicial dos dados na TABELA: ") # Essa é o numero da linha da tabela excel
execelFim = input("Qual é o numero da linha final dos dados na TABELA: ") # Numero final da linha da tabela excel
datas = input("Digite a DATA que voce quer fazer o encerrante: ") # data dos encerrantes
datas2 = input("Digite a DATA2 que voce quer fazer o encerrante: ") # data dos encerrantes
nomeDoTecnico = input("Digite o nome do tecnico: ")
cnpj = input("Digite o cnpj: ")
cpf = input("Digite o cpf: ")

# Esse for e para pegar aquele valor da linha do execel que vamos pular.
opcaoDaSegundaForma = input("Nesse excel possui linhas que precisar ser executada em segunda forma? (S/N): ").upper()

if(opcaoDaSegundaForma == 'S'):
    quantLinhas = int(input("Total de linhas serao ignoradas: "))
    for i in range(quantLinhas):
        linhaQueVaiSerExecutadoEmSegundaForma = int(input(f"Digite o  {i+1} número da linha do execel que vai ser ignorada: "))
        lista_linhas_Excel_Segunda_Forma.append(aba.cell(row=linhaQueVaiSerExecutadoEmSegundaForma, column=1).value)
print(lista_linhas_Excel_Segunda_Forma)
# Fim da lista das linhas ignoradas ================================

for linha in aba['A'+execelInicio:'A'+execelFim]:
    celula = linha[0]
    lista_bicos.append(celula.value)
print("Bicos:", lista_bicos)

for linha in aba['B'+execelInicio:'B'+execelFim]:
    celula = linha[0]
    lista_nomes.append(celula.value)
print("Nomes: ", lista_nomes)

for linha in aba['F'+execelInicio:'F'+execelFim]:
    celula = linha[0]
    lista_diferenca.append(celula.value)
print("Diferença:", lista_diferenca)

#Aqui estou juntando dois array lista_bicos e lista_nomes formando isso por exemplo: 10-GASOLINA COMUM e coloando isso na lista (lista_bicos_nomes)
for item_x,item_y in zip(lista_bicos,lista_nomes):
    combinacao = f"{item_x}-{item_y}"
    lista_bicos_nomes.append(combinacao)
print(lista_bicos_nomes)
#Fim da junção das duas lista e formando outra lista ========================

bicoAtual = 0 #padrao é 0
input("Pronto para iniciar a automação ? (Dê Enter aqui)") # Uma trava aqui só para verificar se está tudo nos conformes (Caso não precise mais é só tirar)

# Automação começa aqui ============================
pyautogui.moveTo(pos.x1,pos.y1,1.00)
pyautogui.click()
pyautogui.moveTo(pos.x2,pos.y2,1.00)
pyautogui.click()

while (bicoAtual < int(lista_bicos[-1])):
    pyautogui.click()
    time.sleep(1)
    if lista_bicos[bicoAtual] in lista_linhas_Excel_Segunda_Forma:
        pyautogui.write(lista_bicos[bicoAtual]+"-")
        pyautogui.press('end')
    else:
        pyautogui.write(lista_bicos_nomes[bicoAtual],1.00)
    time.sleep(1)
    pyautogui.press('tab')
    time.sleep(1)
    pyautogui.write(datas)
    time.sleep(1)
    pyautogui.press('tab')
    time.sleep(3)
    pyautogui.write(datas2)
    time.sleep(1)
    pyautogui.press('tab')
    time.sleep(1)
    pyautogui.press('space')
    time.sleep(1)
    pyautogui.moveTo(pos.x3,pos.y3)
    time.sleep(1)
    pyautogui.click()
    time.sleep(1)
    pyautogui.moveTo(pos.x4,pos.y4)
    time.sleep(1)
    pyautogui.click()
    time.sleep(1)
    pyautogui.moveTo(pos.x5,pos.y5)
    time.sleep(1)
    pyautogui.click()
    time.sleep(1)
    pyautogui.hotkey('ctrl','c')
    time.sleep(1)
    encerranteInicial = pyperclip.paste().replace(",",".")
    time.sleep(1)
    valorCelula = Decimal(lista_diferenca[bicoAtual])
    valorCelulaFormatado = valorCelula.quantize(Decimal('0.000'),rounding=ROUND_HALF_UP)
    consertoDoEncerrante = valorCelulaFormatado + Decimal(encerranteInicial)
    formatadoConsertoDoEncerrante = consertoDoEncerrante.quantize(Decimal('0.000'),rounding=ROUND_HALF_UP)
    time.sleep(1)
    pyautogui.write(str(formatadoConsertoDoEncerrante))
    time.sleep(1)
    pyautogui.click()
    time.sleep(2)
    pyautogui.press('tab')
    time.sleep(1)
    pyautogui.write("AJUSTE ENCERRANTE")
    time.sleep(1)
    pyautogui.press('tab')
    time.sleep(1)
    pyautogui.write('1')
    time.sleep(1)
    pyautogui.press('tab')
    time.sleep(1)
    pyautogui.write(nomeDoTecnico)
    time.sleep(1)
    pyautogui.press('tab')
    time.sleep(1)
    pyautogui.write(cnpj)
    time.sleep(1)
    pyautogui.press('tab')
    time.sleep(1)
    pyautogui.write(cpf)
    pyautogui.press('tab')
    time.sleep(1)
    pyautogui.press('space')
    time.sleep(1)
    pyautogui.press('space')
    time.sleep(1)
    pyautogui.press('tab')
    time.sleep(1)
    pyautogui.moveTo(pos.x1,pos.y1,1.00)
    pyautogui.moveTo(pos.x2,pos.y2,1.00)
    time.sleep(5)
    bicoAtual += 1